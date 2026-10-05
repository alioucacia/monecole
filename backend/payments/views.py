import csv
from datetime import date, datetime
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation
from io import BytesIO

import openpyxl
from openpyxl.utils import get_column_letter
from django.db import transaction
from django.db.models import ProtectedError, Sum
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.shortcuts import get_object_or_404
from django.template.loader import render_to_string
from django.utils import timezone
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import ValidationError
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView
from xhtml2pdf import pisa

from academics.annee import AnneeScolaireFilterBackend, annee_courante, filtre_eleves_annee
from accounts.permissions import IsAdmin, IsAdminOrComptabilite, IsAdminOrComptabiliteOrReadOnly
from people.views import _image_data_uri, _mm_px

from .models import MOIS_MENSUALITE, PERIODICITES_SCOLARITE, est_scolarite, USAGES_INSCRIPTION, CategorieDepense, Depense, Frais, Paiement, TarifClasse, TypeFrais
from .serializers import CategorieDepenseSerializer, DepenseSerializer, FraisSerializer, PaiementSerializer, TarifClasseSerializer, TypeFraisSerializer


def _mois_entre(date_debut, date_fin):
    """Liste des premiers jours de mois entre `date_debut` et `date_fin` (inclus)."""
    mois, courant = [], date_debut.replace(day=1)
    fin = date_fin.replace(day=1)
    while courant <= fin:
        mois.append(courant)
        courant = date(courant.year + 1, 1, 1) if courant.month == 12 else date(courant.year, courant.month + 1, 1)
    return mois


# Mois de mensualité : voir payments.models.MOIS_MENSUALITE (Octobre → Juin). Dans le suivi
# mensuel, la colonne Inscription/Réinscription remplace Septembre.
# Mois couverts par chaque tranche de scolarité (frais de périodicité « Tranche »), dans l'ordre
# des Periode de l'année (Trimestre 1, 2, 3) : la 1re tranche couvre aussi Juin, dernier mois.
MOIS_PAR_TRANCHE = ((10, 11, 12, 6), (1, 2, 3), (4, 5))


def _mois_mensualite(annee_scolaire):
    """Premiers jours des mois de mensualité (Octobre → Juin) de l'année scolaire."""
    return [
        m for m in _mois_entre(annee_scolaire.date_debut, annee_scolaire.date_fin)
        if m.month in MOIS_MENSUALITE
    ]


def _frais_suivi_par_eleve(eleves, annee_scolaire) -> dict[int, list[Frais]]:
    """Tous les frais de l'année des élèves donnés, paiements préchargés — 2 requêtes au total
    quel que soit le nombre d'élèves (le suivi d'un cycle entier en compte des centaines)."""
    par_eleve: dict[int, list[Frais]] = {e.id: [] for e in eleves}
    frais = (
        Frais.objects.filter(eleve__in=eleves, annee_scolaire=annee_scolaire)
        .select_related("type_frais", "eleve").prefetch_related("paiements").order_by("date_echeance")
    )
    for f in frais:
        par_eleve.setdefault(f.eleve_id, []).append(f)
    return par_eleve


def _total_paye(frais, **filtre) -> Decimal:
    """Total versé sur un frais (paiements préchargés), éventuellement limité à une tranche."""
    periode_id = filtre.get("periode_id")
    return sum(
        (p.montant for p in frais.paiements.all() if periode_id is None or p.periode_id == periode_id),
        Decimal("0"),
    )


def _statut(du, paye) -> str:
    if du <= 0 or paye >= du:
        return "paye"
    return "partiel" if paye > 0 else "non_paye"


def _tarifs_tranches(eleve, annee_scolaire) -> dict[int, Decimal]:
    """Montant de chaque tranche numérotée (types « 1ère Tranche », « 2ème Tranche »... — voir
    TypeFrais.numero_tranche) pour la classe de l'élève : tarif de la classe pour l'année s'il y
    en a un, sinon le montant standard du type."""
    types = [
        t for t in TypeFrais.objects.filter(
            ecole_id=eleve.user.ecole_id, periodicite=TypeFrais.Periodicite.TRIMESTRIEL,
        ).exclude(usage__in=USAGES_INSCRIPTION)
        if t.numero_tranche
    ]
    tarifs_classe = {}
    if eleve.classe_id:
        tarifs_classe = dict(
            TarifClasse.objects.filter(
                classe_id=eleve.classe_id, annee_scolaire=annee_scolaire, type_frais__in=types,
            ).values_list("type_frais_id", "montant")
        )
    tarifs = {}
    for t in sorted(types, key=lambda t: t.id):
        tarifs.setdefault(t.numero_tranche, tarifs_classe.get(t.id, t.montant_standard))
    return tarifs


def _calculer_suivi_mensuel(eleve, annee_scolaire, frais=None, periodes=None, cache_tarifs=None):
    """Statut payé/partiel/non payé de chaque mois de mensualité (Octobre → Juin, passés et à
    venir) de la scolarité d'un élève, quelle que soit sa formule de paiement :

    - Mensuel (type de périodicité « Mensuel ») : montant dû d'un mois = `montant_du` du frais
      dont l'échéance tombe ce mois-là ; un mois dont le frais n'a pas encore été créé reste dû
      au tarif mensuel du frais de ce type le plus proche (sinon il disparaîtrait du suivi au
      lieu d'apparaître « Non payé »). Payé = paiements dont `Paiement.mois` est ce mois (ou,
      données anciennes sans mois, sur le frais dont l'échéance tombe ce mois-là).
    - Annuel (hors inscription) : le frais couvre les 9 mois — dû et versé répartis à parts
      égales, donc un paiement annuel complet affiche les 9 mois payés.
    - Tranche (hors inscription) : chaque tranche (Periode de l'année, dans l'ordre) couvre ses
      mois de MOIS_PAR_TRANCHE — la 1re tranche payée affiche Octobre, Novembre, Décembre et
      Juin payés, la 2e Janvier à Mars, la 3e Avril et Mai.

    `montant_du` tient compte de la catégorie de paiement (Fondation 50 %...) : un élève exonéré
    (Fondation gratuite, inscription seulement) n'a aucun mois à suivre. `a_venir` : mois pas
    encore commencé (non compté dans les impayés). `couvert_par` : « Annuel » / « Tranche N »
    quand le mois est réglé par un frais annuel ou une tranche plutôt que mois par mois.

    Tranches numérotées (types « 1ère Tranche », « 2ème Tranche »... chacun avec son montant) :
    les mois d'une tranche dont le frais n'est pas (encore) créé restent dus au tarif de cette
    tranche — « Non payé » tant qu'elle n'est pas payée, et non « Payé » faute de montant dû.

    `frais` / `periodes` / `cache_tarifs` (tarifs des tranches par classe) : préchargés par
    l'appelant pour le suivi d'une classe entière (voir `_frais_suivi_par_eleve`), sinon chargés
    ici."""
    from grades.models import Periode

    if eleve.exonere_fratrie:
        return []  # « Élève Bonus » : aucune scolarité à suivre (voir people/fratrie.py)
    if frais is None:
        frais = _frais_suivi_par_eleve([eleve], annee_scolaire)[eleve.id]
    hors_inscription = [f for f in frais if f.type_frais.usage not in USAGES_INSCRIPTION]
    mensuels = [f for f in hors_inscription if f.type_frais.est_mensuel]
    annuels = [f for f in hors_inscription if f.type_frais.periodicite == TypeFrais.Periodicite.ANNUEL]
    tranches = [f for f in hors_inscription if f.type_frais.periodicite == TypeFrais.Periodicite.TRIMESTRIEL]
    tous_les_mois = _mois_mensualite(annee_scolaire)
    if not (mensuels or annuels or tranches) or not tous_les_mois:
        return []

    du_par_mois = {mois: Decimal("0") for mois in tous_les_mois}
    paye_par_mois = {mois: Decimal("0") for mois in tous_les_mois}
    couvert_par = {mois: [] for mois in tous_les_mois}

    # 1) Mensuel, type de frais par type de frais.
    par_type: dict[int, list[Frais]] = {}
    for f in mensuels:
        par_type.setdefault(f.type_frais_id, []).append(f)
    for frais_du_type in par_type.values():
        par_mois_echeance = {}
        for f in frais_du_type:
            par_mois_echeance.setdefault(f.mois_reference, []).append(f)
        for mois in tous_les_mois:
            if mois in par_mois_echeance:
                du_par_mois[mois] += sum((f.montant_du for f in par_mois_echeance[mois]), Decimal("0"))
            else:
                reference = min(frais_du_type, key=lambda f: abs((f.mois_reference - mois).days))
                du_par_mois[mois] += reference.montant_du
        for f in frais_du_type:
            for p in f.paiements.all():
                mois = p.mois or f.mois_reference
                if mois in paye_par_mois:
                    paye_par_mois[mois] += p.montant

    def repartir(mois_couverts, du, paye, libelle):
        n = len(mois_couverts)
        for mois in mois_couverts:
            du_par_mois[mois] += (du / n).quantize(Decimal("0.01"))
            paye_par_mois[mois] += (paye / n).quantize(Decimal("0.01"))
            couvert_par[mois].append(libelle)

    # 2) Annuel : les 9 mois.
    for f in annuels:
        repartir(tous_les_mois, f.montant_du, _total_paye(f), "Annuel")

    # 3) Tranches : chaque Periode de l'année couvre ses mois (voir MOIS_PAR_TRANCHE) — seuls
    #    les mois d'une tranche payée passent « Payé », les autres restent « Non payé » jusqu'au
    #    paiement de leur tranche. Un versement sans tranche précisée (données antérieures à
    #    l'obligation de la choisir) complète les tranches dans l'ordre : 1re, puis 2e...
    non_payes_sans_tarif = set()
    if tranches:
        if periodes is None:
            periodes = list(Periode.objects.filter(annee_scolaire=annee_scolaire).order_by("date_debut"))
        periodes_tranches = periodes[:len(MOIS_PAR_TRANCHE)]
        for f in tranches:
            numero = f.type_frais.numero_tranche
            if numero:
                # Type propre à une tranche (« 2ème Tranche », avec son propre montant) : il ne
                # couvre que les mois de SA tranche, et tout ce qui y est versé compte pour elle.
                if numero <= len(MOIS_PAR_TRANCHE):
                    mois_tranche = [m for m in tous_les_mois if m.month in MOIS_PAR_TRANCHE[numero - 1]]
                    if mois_tranche:
                        repartir(mois_tranche, f.montant_du, _total_paye(f), f"Tranche {numero}")
                continue
            sans_tranche = sum(
                (p.montant for p in f.paiements.all() if p.periode_id not in {pe.id for pe in periodes_tranches}),
                Decimal("0"),
            )
            for index, periode in enumerate(periodes_tranches):
                verse = _total_paye(f, periode_id=periode.id)
                complement = min(sans_tranche, max(f.montant_du - verse, Decimal("0")))
                sans_tranche -= complement
                mois_tranche = [m for m in tous_les_mois if m.month in MOIS_PAR_TRANCHE[index]]
                if mois_tranche:
                    repartir(mois_tranche, f.montant_du, verse + complement, f"Tranche {index + 1}")

        # Tranches numérotées dont l'élève n'a pas encore de frais : dues à leur tarif.
        numeros_presents = {f.type_frais.numero_tranche for f in tranches if f.type_frais.numero_tranche}
        if numeros_presents:
            cle = (eleve.classe_id, annee_scolaire.id)
            if cache_tarifs is not None and cle in cache_tarifs:
                tarifs = cache_tarifs[cle]
            else:
                tarifs = _tarifs_tranches(eleve, annee_scolaire)
                if cache_tarifs is not None:
                    cache_tarifs[cle] = tarifs
            for numero in range(1, len(MOIS_PAR_TRANCHE) + 1):
                if numero in numeros_presents:
                    continue
                mois_tranche = [m for m in tous_les_mois if m.month in MOIS_PAR_TRANCHE[numero - 1]]
                if not mois_tranche:
                    continue
                if tarifs.get(numero):
                    repartir(mois_tranche, tarifs[numero], Decimal("0"), f"Tranche {numero}")
                else:
                    # Aucune « N-ième Tranche » paramétrée : montant inconnu, mais rien n'est payé.
                    non_payes_sans_tarif.update(mois_tranche)

    if all(du <= 0 for du in du_par_mois.values()):
        return []  # élève exonéré de mensualité : rien à suivre

    mois_courant = date.today().replace(day=1)
    resultat = []
    for mois in tous_les_mois:
        du, paye = du_par_mois[mois], paye_par_mois[mois]
        resultat.append({
            "mois": mois.strftime("%Y-%m"), "montant_du": du, "montant_paye": paye,
            "reste": max(du - paye, Decimal("0")),
            "statut": "non_paye" if mois in non_payes_sans_tarif and paye <= 0 else _statut(du, paye),
            "a_venir": mois > mois_courant,
            "couvert_par": ", ".join(dict.fromkeys(couvert_par[mois])),
        })
    return resultat


def _tarifs_inscription(eleve, annee_scolaire) -> dict[str, Decimal]:
    """Montant de l'inscription et de la réinscription (TypeFrais.usage) pour la classe de
    l'élève : tarif de la classe pour l'année s'il y en a un, sinon le montant standard du type."""
    tarifs = {}
    for usage in USAGES_INSCRIPTION:
        type_frais = TypeFrais.objects.filter(ecole_id=eleve.user.ecole_id, usage=usage).first()
        if not type_frais:
            continue
        tarif = None
        if eleve.classe_id:
            tarif = TarifClasse.objects.filter(
                type_frais=type_frais, classe_id=eleve.classe_id, annee_scolaire=annee_scolaire,
            ).values_list("montant", flat=True).first()
        tarifs[usage] = tarif if tarif is not None else type_frais.montant_standard
    return tarifs


def _suivi_inscription(frais, eleve=None, annee_scolaire=None, cache_tarifs=None):
    """Frais d'inscription ou de réinscription de l'année (colonne qui remplace Septembre dans le
    suivi mensuel). Si l'élève n'en a pas encore (frais pas créé), il reste dû — « Non payé » —
    au tarif de sa réinscription (élève au statut « Réinscription ») ou de son inscription ;
    `None` seulement sans élève/année pour le calculer."""
    frais_inscription = [f for f in frais if f.type_frais.usage in USAGES_INSCRIPTION]
    if not frais_inscription:
        if eleve is None or annee_scolaire is None:
            return None
        cle = ("inscription", eleve.classe_id, annee_scolaire.id)
        if cache_tarifs is not None and cle in cache_tarifs:
            tarifs = cache_tarifs[cle]
        else:
            tarifs = _tarifs_inscription(eleve, annee_scolaire)
            if cache_tarifs is not None:
                cache_tarifs[cle] = tarifs
        reinscription = eleve.statut_inscription == eleve.StatutInscription.REINSCRIPTION
        usage = TypeFrais.Usage.REINSCRIPTION if reinscription else TypeFrais.Usage.INSCRIPTION
        tarif = tarifs.get(usage)
        if tarif is None:  # type de ce statut non paramétré : on se rabat sur l'autre
            tarif = next(iter(tarifs.values()), None)
        du = (
            (tarif * eleve.facteur_inscription_reinscription).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
            if tarif is not None else Decimal("0")
        )
        return {
            "libelle": "Réinscription" if reinscription else "Inscription",
            "montant_du": du, "montant_paye": Decimal("0"), "reste": du,
            # Aucun montant paramétré : rien n'est payé pour autant. Exonéré (Fondation 100 %) :
            # rien à payer.
            "statut": "non_paye" if tarif is None or du > 0 else "paye",
        }
    du = sum((f.montant_du for f in frais_inscription), Decimal("0"))
    paye = sum((_total_paye(f) for f in frais_inscription), Decimal("0"))
    reinscription = any(f.type_frais.usage == TypeFrais.Usage.REINSCRIPTION for f in frais_inscription)
    return {
        "libelle": "Réinscription" if reinscription else "Inscription",
        "montant_du": du, "montant_paye": paye, "reste": max(du - paye, Decimal("0")), "statut": _statut(du, paye),
    }


def _situation_paiement(eleve_suivi, mois_filtre="") -> str | None:
    """« payes » si tout ce qui est dû est réglé, « non_payes » sinon (impayé ou partiel) — sur
    le mois filtré, ou à défaut sur l'inscription/réinscription et les mois déjà commencés (un
    mois à venir n'est pas encore dû). `None` : rien à suivre pour cet élève."""
    if mois_filtre:
        colonnes = [m for m in eleve_suivi["mois"] if m["mois"] == mois_filtre]
    else:
        colonnes = [m for m in eleve_suivi["mois"] if not m["a_venir"]]
        if eleve_suivi["inscription"]:
            colonnes.append(eleve_suivi["inscription"])
    if not colonnes:
        return None
    return "payes" if all(c["statut"] == "paye" for c in colonnes) else "non_payes"


def _fiche_context(frais: Frais) -> dict:
    return {
        "eleve_nom": frais.eleve.user.get_full_name(),
        "matricule": frais.eleve.matricule,
        "classe": frais.eleve.classe.nom if frais.eleve.classe else None,
        "type_frais": frais.type_frais.nom,
        "annee_scolaire": frais.annee_scolaire.libelle,
        "montant": frais.montant,
        "montant_du": frais.montant_du,
        "montant_paye": frais.montant_paye,
        "solde": frais.solde,
        "statut": frais.statut,
        "date_echeance": frais.date_echeance,
        "paiements": list(frais.paiements.select_related("enregistre_par").order_by("date_paiement")),
    }


def _render_fiches_pdf(fiches: list[dict], titre: str, ecole=None) -> bytes:
    """`ecole` sert à personnaliser le document (logo, couleurs, modèle choisis par le Super
    Admin — voir `Ecole.modele_recu`/`couleur_principale`/`couleur_secondaire`) : toutes les
    fiches d'un même appel appartiennent à la même école (déjà filtrées par `get_queryset()`),
    donc une seule résolution suffit plutôt que de la refaire fiche par fiche."""
    html = render_to_string("payments/fiche_paiement_pdf.html", {
        "fiches": fiches, "titre": titre,
        "ecole_nom": ecole.nom if ecole else "Taly-School",
        "ecole_logo_data_uri": _image_data_uri(ecole.logo, _mm_px(11, 11), mode="contain") if ecole else None,
        "couleur_principale": ecole.couleur_principale if ecole else "#14304f",
        "couleur_secondaire": ecole.couleur_secondaire if ecole else "#b8860b",
        "modele": ecole.modele_recu if ecole else 1,
    })
    buffer = BytesIO()
    pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
    return buffer.getvalue()


class TypeFraisViewSet(viewsets.ModelViewSet):
    queryset = TypeFrais.objects.all()
    serializer_class = TypeFraisSerializer
    permission_classes = [IsAdminOrComptabiliteOrReadOnly]

    def get_queryset(self):
        return super().get_queryset().filter(ecole_id=self.request.user.ecole_id)

    def perform_create(self, serializer):
        serializer.save(ecole=self.request.user.ecole)

    def perform_destroy(self, instance):
        # type_frais est protégé (on_delete=PROTECT) tant que des Frais y sont rattachés,
        # pour ne jamais perdre l'historique des paiements déjà encaissés.
        try:
            instance.delete()
        except ProtectedError:
            raise ValidationError(
                "Impossible de supprimer ce type de frais : des frais y sont déjà rattachés "
                "(historique de paiements). Vous pouvez le renommer à la place."
            )

    @action(detail=False, methods=["get"], url_path="tarif-par-usage")
    def tarif_par_usage(self, request):
        """Résout en un seul appel le montant d'inscription ou de réinscription configuré pour
        une classe — combine `TypeFrais.usage` (lequel des types de frais de l'école EST le frais
        d'inscription/réinscription, voir ce champ) et `TarifClasse` (son montant POUR CETTE
        classe) : évite au frontend (StudentsPage à la création d'un élève, ReinscriptionPage) de
        connaître l'id du type de frais ou de faire deux requêtes. Renvoie systématiquement 200,
        avec `type_frais`/`montant` à `null` si rien n'est configuré (l'admin n'a pas encore créé
        ce type de frais, ou pas encore réglé son tarif pour cette classe) — laisser le frontend
        décider comment l'indiquer plutôt que de renvoyer une erreur pour un cas de configuration
        parfaitement normal (nouvelle école, pas encore paramétrée)."""
        from academics.models import Classe

        usage = request.query_params.get("usage")
        classe_id = request.query_params.get("classe")
        if usage not in (TypeFrais.Usage.INSCRIPTION, TypeFrais.Usage.REINSCRIPTION):
            raise ValidationError("Le paramètre 'usage' doit valoir 'inscription' ou 'reinscription'.")
        if not classe_id:
            raise ValidationError("Le paramètre 'classe' est requis.")
        classe = get_object_or_404(Classe, pk=classe_id, annee_scolaire__ecole_id=request.user.ecole_id)

        type_frais = TypeFrais.objects.filter(ecole_id=request.user.ecole_id, usage=usage).first()
        if not type_frais:
            return Response({"type_frais": None, "type_frais_nom": None, "montant": None})

        tarif = TarifClasse.objects.filter(
            ecole_id=request.user.ecole_id, type_frais=type_frais,
            classe=classe, annee_scolaire_id=classe.annee_scolaire_id,
        ).first()
        return Response({
            "type_frais": type_frais.id,
            "type_frais_nom": type_frais.nom,
            # Faute d'un tarif spécifique à cette classe, le tarif standard du type de frais sert
            # de repli — cohérent avec le reste de l'app (FraisViewSet.generer_pour_classe fait
            # de même) plutôt que de renvoyer `null` et laisser croire qu'il n'y a rien à payer.
            "montant": str(tarif.montant) if tarif else str(type_frais.montant_standard),
            "montant_specifique_classe": tarif is not None,
        })


class TarifClasseViewSet(viewsets.ModelViewSet):
    """Paramétrage, par classe et par année scolaire, du montant de chaque type de frais —
    ex: la scolarité peut coûter plus cher en Terminale qu'en 6ème. Sert de valeur par défaut
    lors de la création des frais d'un élève (voir `FraisViewSet.generer_pour_classe`)."""

    queryset = TarifClasse.objects.select_related("type_frais", "classe", "annee_scolaire")
    serializer_class = TarifClasseSerializer
    permission_classes = [IsAdminOrComptabilite]
    filterset_fields = {
        "annee_scolaire": ["exact"],
        "classe": ["exact"],
        "classe__niveau": ["exact"],
        "type_frais": ["exact"],
    }
    annee_scolaire_field = "annee_scolaire"  # listes limitées à l'année affichée (academics/annee.py)

    def get_queryset(self):
        return super().get_queryset().filter(ecole_id=self.request.user.ecole_id)

    def perform_create(self, serializer):
        serializer.save(ecole=self.request.user.ecole)

    @action(detail=False, methods=["post"], url_path="definir")
    def definir(self, request):
        """Crée ou met à jour (upsert) le tarif d'un type de frais pour une classe/année — utilisé
        par la grille de paramétrage (une case = un tarif), pour ne pas avoir à distinguer
        création/modification côté frontend."""
        from academics.models import AnneeScolaire, Classe

        type_frais_id = request.data.get("type_frais")
        classe_id = request.data.get("classe")
        annee_id = request.data.get("annee_scolaire")
        montant = request.data.get("montant")
        if not (type_frais_id and classe_id and annee_id and montant not in (None, "")):
            raise ValidationError("Les champs 'type_frais', 'classe', 'annee_scolaire' et 'montant' sont requis.")

        type_frais = get_object_or_404(TypeFrais, pk=type_frais_id, ecole_id=request.user.ecole_id)
        classe = get_object_or_404(Classe, pk=classe_id, annee_scolaire__ecole_id=request.user.ecole_id)
        annee = get_object_or_404(AnneeScolaire, pk=annee_id, ecole_id=request.user.ecole_id)

        tarif, _ = TarifClasse.objects.update_or_create(
            ecole=request.user.ecole, type_frais=type_frais, classe=classe, annee_scolaire=annee,
            defaults={"montant": montant},
        )
        return Response(TarifClasseSerializer(tarif).data)


class CategorieDepenseViewSet(viewsets.ModelViewSet):
    """Catégories de dépense de l'établissement — librement gérées par son administrateur
    (renommer, ajouter, supprimer), contrairement aux choix fixes qu'elles remplacent. Un jeu de
    départ (CATEGORIES_DEPENSE_PAR_DEFAUT) est créé automatiquement à la création de l'école."""

    queryset = CategorieDepense.objects.all()
    serializer_class = CategorieDepenseSerializer
    permission_classes = [IsAdminOrComptabilite]

    def get_queryset(self):
        return super().get_queryset().filter(ecole_id=self.request.user.ecole_id)

    def perform_create(self, serializer):
        serializer.save(ecole=self.request.user.ecole)

    def perform_destroy(self, instance):
        # PROTECT sur Depense.categorie : une catégorie déjà utilisée ne doit pas pouvoir
        # disparaître en laissant des dépenses orphelines/sans catégorie.
        try:
            instance.delete()
        except ProtectedError:
            raise ValidationError(
                f"Impossible de supprimer « {instance.nom} » : des dépenses y sont déjà rattachées. "
                "Vous pouvez la renommer à la place."
            )


class DepenseViewSet(viewsets.ModelViewSet):
    """Sorties de caisse de l'établissement (hors salaires enseignants) — le pendant de
    `FraisViewSet`/`PaiementViewSet` côté dépenses, alimente le tableau de bord Caisse
    (voir `CaisseView`) et le rapport journalier (simple filtre de date sur cette même vue)."""

    queryset = Depense.objects.select_related("enregistre_par", "categorie")
    serializer_class = DepenseSerializer
    permission_classes = [IsAdminOrComptabilite]
    filterset_fields = {
        "date": ["exact", "gte", "lte"],
        "categorie": ["exact"],
        "mode_paiement": ["exact"],
    }
    annee_date_field = "date"

    def get_queryset(self):
        return super().get_queryset().filter(ecole_id=self.request.user.ecole_id)

    def perform_create(self, serializer):
        serializer.save(ecole=self.request.user.ecole, enregistre_par=self.request.user)

    @action(detail=False, methods=["get"], url_path="summary")
    def summary(self, request):
        qs = self.filter_queryset(self.get_queryset())
        total = qs.aggregate(total=Sum("montant"))["total"] or Decimal("0")
        return Response({"total": total, "nombre": qs.count()})

    @action(detail=False, methods=["get"], url_path="export")
    def export(self, request):
        """Export CSV des dépenses (respecte les filtres de la liste, ex: ?date__gte=…)."""
        queryset = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
        response["Content-Disposition"] = 'attachment; filename="depenses.csv"'
        writer = csv.writer(response, delimiter=";")
        writer.writerow(["Date", "Catégorie", "Motif", "Montant", "Mode de paiement", "Référence", "Responsable", "Enregistré par"])
        for depense in queryset:
            writer.writerow([
                depense.date, depense.categorie.nom, depense.motif, depense.montant,
                depense.get_mode_paiement_display(), depense.reference, depense.responsable,
                depense.enregistre_par.get_full_name() if depense.enregistre_par else "",
            ])
        return response


class CaisseView(APIView):
    """Tableau de bord Caisse : combine les rentrées (paiements élèves) et les sorties (dépenses)
    de l'école sur une période — par défaut aujourd'hui, ce qui couvre directement le rapport
    journalier. `?date_debut=&date_fin=` (YYYY-MM-DD) pour une période personnalisée."""

    permission_classes = [IsAuthenticated, IsAdminOrComptabilite]

    def _periode(self, request):
        aujourdhui = date.today()
        date_debut = request.query_params.get("date_debut")
        date_fin = request.query_params.get("date_fin")
        try:
            debut = date.fromisoformat(date_debut) if date_debut else aujourdhui
            fin = date.fromisoformat(date_fin) if date_fin else aujourdhui
        except ValueError:
            raise ValidationError("Format de date invalide — attendu AAAA-MM-JJ.")
        if debut > fin:
            raise ValidationError("La date de début doit précéder la date de fin.")
        return debut, fin

    def _donnees(self, request):
        """Retourne déjà les rentrées/sorties sous forme de dicts sérialisés (pas les objets
        modèle bruts) : cette même forme sert à la fois à la réponse JSON (`CaisseView.get`) et
        au contexte du template PDF (`CaissePdfView.get`), pour que les deux ne puissent jamais
        diverger — un objet `Paiement`/`Depense` brut n'a pas d'attribut `eleve_nom`/
        `mode_paiement_display` et rendrait ces colonnes silencieusement vides dans le PDF."""
        ecole_id = request.user.ecole_id
        debut, fin = self._periode(request)
        paiements = (
            Paiement.objects.filter(
                frais__eleve__user__ecole_id=ecole_id, date_paiement__gte=debut, date_paiement__lte=fin,
            ).select_related("frais__eleve__user", "frais__type_frais", "enregistre_par").order_by("date_paiement")
        )
        depenses = (
            Depense.objects.filter(ecole_id=ecole_id, date__gte=debut, date__lte=fin)
            .select_related("enregistre_par", "categorie").order_by("date")
        )
        rentrees = [
            {
                "id": p.id, "date": p.date_paiement, "eleve_nom": p.frais.eleve.user.get_full_name(),
                "type_frais_nom": p.frais.type_frais.nom, "montant": p.montant,
                "mode_paiement": p.mode_paiement, "mode_paiement_display": p.get_mode_paiement_display(),
                "reference": p.reference,
                "enregistre_par_nom": p.enregistre_par.get_full_name() if p.enregistre_par else None,
            }
            for p in paiements
        ]
        sorties = [
            {
                "id": d.id, "date": d.date, "categorie": d.categorie_id,
                "categorie_nom": d.categorie.nom, "motif": d.motif, "montant": d.montant,
                "mode_paiement": d.mode_paiement, "mode_paiement_display": d.get_mode_paiement_display(),
                "reference": d.reference, "responsable": d.responsable,
                "enregistre_par_nom": d.enregistre_par.get_full_name() if d.enregistre_par else None,
            }
            for d in depenses
        ]
        total_rentrees = sum((r["montant"] for r in rentrees), Decimal("0"))
        total_sorties = sum((s["montant"] for s in sorties), Decimal("0"))
        return {
            "date_debut": debut, "date_fin": fin,
            "total_rentrees": total_rentrees, "total_sorties": total_sorties,
            "solde": total_rentrees - total_sorties,
            "rentrees": rentrees, "sorties": sorties,
        }

    def get(self, request):
        return Response(self._donnees(request))


class CaissePdfView(APIView):
    """Version imprimable (PDF) du tableau de bord Caisse — le rapport journalier demandé quand
    date_debut = date_fin = aujourd'hui (le réglage par défaut)."""

    permission_classes = [IsAuthenticated, IsAdminOrComptabilite]

    def get(self, request):
        vue = CaisseView()
        donnees = vue._donnees(request)
        ecole = request.user.ecole
        html = render_to_string("payments/caisse_pdf.html", {
            **donnees,
            "ecole_nom": ecole.nom if ecole else "Taly-School",
            "ecole_adresse": ecole.adresse if ecole else "",
            "ecole_telephone": ecole.telephone if ecole else "",
            "ecole_logo_data_uri": _image_data_uri(ecole.logo, _mm_px(16, 16), mode="contain") if ecole and ecole.logo else None,
            "couleur_principale": ecole.couleur_principale if ecole else "#14304f",
            "couleur_secondaire": ecole.couleur_secondaire if ecole else "#b8860b",
            "genere_par": request.user.get_full_name(),
        })
        buffer = BytesIO()
        pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        nom_fichier = f"rapport_caisse_{donnees['date_debut']}_{donnees['date_fin']}.pdf"
        response["Content-Disposition"] = f'attachment; filename="{nom_fichier}"'
        return response


class FraisViewSet(viewsets.ModelViewSet):
    queryset = Frais.objects.select_related("eleve__user", "type_frais", "annee_scolaire").prefetch_related("paiements")
    serializer_class = FraisSerializer
    permission_classes = [IsAdminOrComptabilite]
    filterset_fields = {
        "eleve": ["exact"],
        "type_frais": ["exact"],
        "annee_scolaire": ["exact"],
        "eleve__classe": ["exact"],
        "eleve__classe__niveau": ["exact"],
        "eleve__classe__cycle": ["exact"],
    }
    annee_scolaire_field = "annee_scolaire"
    # Recherche libre par nom/prénom/matricule de l'élève (onglet Paiements — filtre par
    # classe/nom/niveau demandé par l'établissement).
    search_fields = ["eleve__user__first_name", "eleve__user__last_name", "eleve__matricule"]

    def get_queryset(self):
        qs = super().get_queryset().filter(eleve__user__ecole_id=self.request.user.ecole_id)
        user = self.request.user
        if user.role == "student" and hasattr(user, "eleve_profile"):
            return qs.filter(eleve=user.eleve_profile)
        if user.role == "parent":
            return qs.filter(eleve__parent=user)
        return qs

    def filter_queryset(self, queryset):
        """`?masquer_payes=1` : exclut les frais au statut « Payé » (liste Paiements). Le statut
        dépend de `montant_du` (réduction propre à l'élève), calculé en Python : total versé
        annoté en une requête, puis tri des frais restant dus."""
        queryset = super().filter_queryset(queryset)
        if self.request.query_params.get("masquer_payes") not in ("1", "true"):
            return queryset
        annotes = (
            queryset.select_related("eleve", "type_frais").prefetch_related(None)
            .annotate(total_verse=Coalesce(Sum("paiements__montant"), Decimal("0")))
        )
        ids_dus = [f.id for f in annotes if f.montant_du > 0 and f.total_verse < f.montant_du]
        return queryset.filter(id__in=ids_dus)

    def get_permissions(self):
        if self.action in ("list", "retrieve", "fiche_paiement", "suivi_mensuel", "proforma"):
            from rest_framework.permissions import IsAuthenticated
            return [IsAuthenticated()]
        return super().get_permissions()

    def perform_destroy(self, instance):
        # La suppression d'un frais est en CASCADE sur ses paiements (voir Paiement.frais) : un
        # frais déjà payé, même partiellement, ne doit donc jamais être supprimé, sous peine de
        # perdre l'historique d'encaissement. Seul un frais encore intégralement impayé (aucun
        # paiement enregistré dessus) peut l'être — ex: un frais généré par erreur.
        if instance.montant_paye > 0:
            raise ValidationError(
                "Impossible de supprimer ce frais : des paiements y sont déjà enregistrés "
                "(historique de paiements). Seul un frais impayé peut être supprimé."
            )
        instance.delete()

    @action(detail=False, methods=["get"], url_path="import-excel-modele", permission_classes=[IsAdmin])
    def import_excel_modele(self, request):
        """Modèle Excel (.xlsx) à remplir pour l'import en masse de frais/paiements historiques
        (reprise de données depuis un autre logiciel de gestion scolaire) — colonnes attendues
        par `import_excel` ci-dessous."""
        classeur = openpyxl.Workbook()
        feuille = classeur.active
        feuille.title = "Frais et paiements"
        entetes = [
            "Matricule élève*", "Type de frais*", "Année scolaire*", "Montant du frais (GNF)*",
            "Date d'échéance (JJ/MM/AAAA)", "Montant déjà payé (GNF)", "Date du paiement (JJ/MM/AAAA)",
            "Mode de paiement (especes/cheque/virement/mobile_money)", "Mois couvert (MM/AAAA)", "Référence",
        ]
        feuille.append(entetes)
        feuille.append([
            "DI2401", "Mensualité", "2024-2025", "500000", "05/10/2024",
            "500000", "03/10/2024", "especes", "10/2024", "",
        ])
        for i, entete in enumerate(entetes, start=1):
            feuille.column_dimensions[get_column_letter(i)].width = max(len(entete) * 0.9, 18)

        buffer = BytesIO()
        classeur.save(buffer)
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="modele_import_frais_paiements.xlsx"'
        return response

    @action(detail=False, methods=["post"], url_path="import-excel", permission_classes=[IsAdmin])
    def import_excel(self, request):
        """Import en masse de frais et de leur historique de paiement depuis un fichier Excel
        (.xlsx, voir `import_excel_modele` pour le format attendu) — pensé pour reprendre les
        données d'un autre logiciel de gestion scolaire sans tout ressaisir à la main.

        Une ligne = un frais (créé s'il n'existe pas encore pour cet élève/type/année, sinon
        réutilisé) + optionnellement UN paiement historique dessus. Pour un frais réglé en
        plusieurs fois par le passé, ajouter une ligne par versement avec le même élève/type/
        année : elles se rattachent toutes au même frais plutôt que d'en recréer un à chaque fois.

        `Paiement.date_paiement` est `auto_now_add` (toujours "aujourd'hui" à la création, voir
        Paiement.save) : impossible d'y écrire la vraie date historique via un simple `.save()`
        — d'où le correctif via `.update()` juste après (le seul moyen de contourner `auto_now_add`).

        Ne s'arrête jamais à la première erreur : chaque ligne est traitée indépendamment, le
        rapport final liste les lignes créées et celles en échec avec leur motif, pour corriger
        et réimporter seulement les lignes en erreur (même logique que
        `EleveProfileViewSet.import_excel`)."""
        from academics.models import AnneeScolaire
        from people.models import EleveProfile

        fichier = request.FILES.get("fichier")
        if not fichier:
            raise ValidationError("Le paramètre 'fichier' (fichier .xlsx) est requis.")
        try:
            classeur = openpyxl.load_workbook(fichier, data_only=True)
        except Exception:
            raise ValidationError("Fichier illisible — vérifiez qu'il s'agit bien d'un fichier Excel (.xlsx) valide.")
        feuille = classeur.active

        ecole_id = request.user.ecole_id
        eleves_par_matricule = {
            e.matricule.strip().lower(): e
            for e in EleveProfile.objects.filter(user__ecole_id=ecole_id).select_related("user")
        }
        types_par_nom = {t.nom.strip().lower(): t for t in TypeFrais.objects.filter(ecole_id=ecole_id)}
        annees_par_libelle = {a.libelle.strip().lower(): a for a in AnneeScolaire.objects.filter(ecole_id=ecole_id)}
        modes_valides = {valeur for valeur, _ in Paiement.ModePaiement.choices}

        def parse_date(valeur):
            """(date, message_erreur) — message_erreur non None seulement si `valeur` est
            renseignée mais illisible ; une valeur vide renvoie (None, None), pas une erreur."""
            if not valeur:
                return None, None
            if hasattr(valeur, "date"):
                return valeur.date(), None
            try:
                return datetime.strptime(str(valeur).strip(), "%d/%m/%Y").date(), None
            except ValueError:
                return None, f"« {valeur} » (format attendu JJ/MM/AAAA)"

        def parse_montant(valeur):
            if valeur in (None, ""):
                return None
            try:
                montant = Decimal(str(valeur).strip().replace(" ", "").replace(",", "."))
                return montant if montant > 0 else None
            except InvalidOperation:
                return None

        lignes = list(feuille.iter_rows(min_row=2, values_only=True))
        frais_crees, paiements_crees = 0, 0
        erreurs = []

        for num_ligne, ligne in enumerate(lignes, start=2):
            if not ligne or all(valeur in (None, "") for valeur in ligne):
                continue  # ligne vide (souvent en fin de feuille) — ignorée silencieusement

            valeurs = (list(ligne) + [None] * 10)[:10]
            (matricule, type_nom, annee_libelle, montant_frais, echeance_brute,
             montant_paye, date_paiement_brute, mode_paiement, mois_brut, reference) = valeurs

            eleve = eleves_par_matricule.get(str(matricule or "").strip().lower())
            if not eleve:
                erreurs.append({"ligne": num_ligne, "message": f"Élève introuvable pour le matricule « {matricule} »."})
                continue

            type_frais = types_par_nom.get(str(type_nom or "").strip().lower())
            if not type_frais:
                erreurs.append({
                    "ligne": num_ligne,
                    "message": f"Type de frais « {type_nom} » introuvable — créez-le d'abord dans Paiements → ⚙️ Types de frais.",
                })
                continue

            annee = annees_par_libelle.get(str(annee_libelle or "").strip().lower())
            if not annee:
                erreurs.append({
                    "ligne": num_ligne,
                    "message": f"Année scolaire « {annee_libelle} » introuvable — créez-la d'abord dans Paramètres de l'école.",
                })
                continue

            montant = parse_montant(montant_frais)
            if montant is None:
                erreurs.append({"ligne": num_ligne, "message": f"Montant du frais invalide : « {montant_frais} »."})
                continue

            echeance, err = parse_date(echeance_brute)
            if err:
                erreurs.append({"ligne": num_ligne, "message": f"Date d'échéance invalide : {err}"})
                continue
            echeance = echeance or annee.date_fin

            filtre = Frais.filtre_equivalents(eleve, type_frais, annee, echeance)
            existant = Frais.objects.filter(filtre).first() if filtre is not None else None
            if existant:
                frais, cree = existant, False
            else:
                frais, cree = Frais.objects.get_or_create(
                    eleve=eleve, type_frais=type_frais, annee_scolaire=annee,
                    defaults={"montant": montant, "date_echeance": echeance},
                )
            if cree:
                frais_crees += 1

            montant_verse = parse_montant(montant_paye)
            if montant_verse is None:
                continue  # ligne "frais seul", sans paiement historique à enregistrer

            date_versement, err = parse_date(date_paiement_brute)
            if err:
                erreurs.append({"ligne": num_ligne, "message": f"Date du paiement invalide : {err}"})
                continue
            date_versement = date_versement or echeance

            mode = str(mode_paiement or "").strip().lower()
            if mode not in modes_valides:
                mode = Paiement.ModePaiement.ESPECES

            mois = None
            if mois_brut:
                try:
                    mois = (
                        mois_brut.date().replace(day=1) if hasattr(mois_brut, "date")
                        else datetime.strptime(str(mois_brut).strip(), "%m/%Y").date().replace(day=1)
                    )
                except ValueError:
                    erreurs.append({"ligne": num_ligne, "message": f"Mois couvert invalide : « {mois_brut} » (format attendu MM/AAAA)."})
                    continue

            # Passe par PaiementSerializer (pas une création ORM directe) pour bénéficier des
            # mêmes garde-fous qu'un encaissement normal (pas de double paiement sur un mois/une
            # tranche déjà soldé·e, pas de dépassement du solde restant) — une ligne qui les
            # viole (ex: total historique incohérent) est signalée comme erreur plutôt
            # qu'importée telle quelle.
            serializer = PaiementSerializer(data={
                "frais": frais.id, "montant": str(montant_verse), "mode_paiement": mode,
                "reference": str(reference or "").strip(), "mois": mois.isoformat() if mois else None,
            })
            if not serializer.is_valid():
                premiere = next(iter(serializer.errors.values()))
                erreurs.append({"ligne": num_ligne, "message": str(premiere[0] if isinstance(premiere, list) else premiere)})
                continue
            paiement = serializer.save(enregistre_par=request.user)
            Paiement.objects.filter(pk=paiement.pk).update(date_paiement=date_versement)
            paiements_crees += 1

            # Même gel de réduction qu'un paiement normal (voir PaiementViewSet.perform_create) —
            # sans ce correctif, un import historique laisserait `facteur_applique` à `None`
            # indéfiniment, et la réduction COURANTE de l'élève s'appliquerait rétroactivement à
            # un frais déjà soldé dans l'ancien logiciel.
            facteur_actuel = frais._facteur_reduction_courant()
            if frais.facteur_applique is None and facteur_actuel is not None:
                frais.facteur_applique = facteur_actuel
                frais.save(update_fields=["facteur_applique"])

        return Response({
            "frais_crees": frais_crees, "paiements_crees": paiements_crees,
            "erreurs": erreurs, "total_lignes": len(lignes),
        })

    @action(detail=False, methods=["get"], url_path="suivi-mensuel")
    def suivi_mensuel(self, request):
        """Statut payé/partiel/non payé, mois par mois, des frais mensuels d'un élève
        (ex: scolarité mensuelle) — pour l'année scolaire active, ou celle précisée."""
        from academics.models import AnneeScolaire
        from people.models import EleveProfile

        eleve_id = request.query_params.get("eleve")
        if not eleve_id:
            raise ValidationError("Le paramètre 'eleve' est requis.")
        eleve = get_object_or_404(EleveProfile, pk=eleve_id, user__ecole_id=request.user.ecole_id)

        if request.user.role == "student" and getattr(request.user, "eleve_profile", None) != eleve:
            raise ValidationError("Vous ne pouvez consulter que votre propre suivi de paiement.")
        if request.user.role == "parent" and eleve.parent_id != request.user.id:
            raise ValidationError("Vous ne pouvez consulter que le suivi de paiement de vos enfants.")

        annee_id = request.query_params.get("annee_scolaire")
        if annee_id:
            annee = get_object_or_404(AnneeScolaire, pk=annee_id, ecole_id=request.user.ecole_id)
        else:
            annee = annee_courante(request)
        if not annee:
            raise ValidationError("Aucune année scolaire active pour votre établissement.")

        frais = _frais_suivi_par_eleve([eleve], annee)[eleve.id]
        return Response({
            "eleve_id": eleve.id, "eleve_nom": eleve.user.get_full_name(),
            "categorie_paiement": eleve.categorie_paiement, "categorie_paiement_display": eleve.get_categorie_paiement_display(),
            "exonere_fratrie": eleve.exonere_fratrie,
            "annee_scolaire": annee.libelle, "mois": _calculer_suivi_mensuel(eleve, annee, frais=frais),
            "inscription": _suivi_inscription(frais, eleve, annee),
        })

    def _grille_suivi_mensuel(self, request) -> dict:
        """Grille de suivi mensuel (payé/partiel/non payé) pour tous les élèves d'une classe — ou,
        sans `classe`, de toutes les classes d'un `cycle` (ou de toute l'école, sans filtre) de
        l'année active. Chaque élève porte aussi le statut de son inscription/réinscription
        (colonne qui remplace Septembre). Partagée par l'écran et sa version imprimable."""
        from academics.models import AnneeScolaire, Classe
        from grades.models import Periode
        from people.models import EleveProfile

        classe_id = request.query_params.get("classe")
        cycle = request.query_params.get("cycle")
        if classe_id:
            classe = get_object_or_404(Classe, pk=classe_id, annee_scolaire__ecole_id=request.user.ecole_id)
            annee = classe.annee_scolaire
            eleves_qs = EleveProfile.objects.filter(classe=classe, actif=True)
            libelle = classe.nom
        else:
            # Sans classe : toutes les classes du cycle, ou par défaut tous les élèves de l'année
            # active (page ouverte sans filtre).
            annee = annee_courante(request)
            if not annee:
                raise ValidationError("Aucune année scolaire active pour votre établissement.")
            # Année passée : élèves de l'époque (historique des classes), même partis depuis.
            eleves_qs = (
                EleveProfile.objects.filter(classe__annee_scolaire=annee, actif=True) if annee.active
                else EleveProfile.objects.filter(filtre_eleves_annee(annee)).distinct()
            )
            if cycle:
                eleves_qs = eleves_qs.filter(classe__cycle=cycle)
            libelle = f"Cycle {Classe.Cycle(cycle).label}" if cycle in Classe.Cycle.values else "Toutes les classes"

        eleves = list(
            eleves_qs.select_related("user", "classe")
            .order_by("classe__niveau", "classe__nom", "user__last_name", "user__first_name")
        )
        frais_par_eleve = _frais_suivi_par_eleve(eleves, annee)
        periodes = list(Periode.objects.filter(annee_scolaire=annee).order_by("date_debut"))
        cache_tarifs = {}
        data = [
            {
                "eleve_id": e.id, "eleve_nom": e.user.get_full_name(), "matricule": e.matricule,
                "classe_nom": e.classe.nom if e.classe else "",
                "categorie_paiement": e.categorie_paiement, "categorie_paiement_display": e.get_categorie_paiement_display(),
                "exonere_fratrie": e.exonere_fratrie,
                "mois": _calculer_suivi_mensuel(
                    e, annee, frais=frais_par_eleve[e.id], periodes=periodes, cache_tarifs=cache_tarifs,
                ),
                "inscription": _suivi_inscription(frais_par_eleve[e.id], e, annee, cache_tarifs),
            }
            for e in eleves
        ]
        return {
            "classe": libelle, "annee_scolaire": annee.libelle,
            "mois": [m.strftime("%Y-%m") for m in _mois_mensualite(annee)],
            "eleves": data,
        }

    @action(detail=False, methods=["get"], url_path="suivi-mensuel-classe")
    def suivi_mensuel_classe(self, request):
        """Grille de suivi mensuel — voir `_grille_suivi_mensuel`."""
        return Response(self._grille_suivi_mensuel(request))

    @action(detail=False, methods=["get"], url_path="suivi-mensuel-classe-pdf")
    def suivi_mensuel_classe_pdf(self, request):
        """Version imprimable (PDF, paysage) du suivi mensuel, avec les mêmes filtres que l'écran :
        `classe` / `cycle`, et `mois` (« AAAA-MM ») pour n'imprimer que ce mois — sans la colonne
        Inscription/Réinscription, comme à l'écran."""
        grille = self._grille_suivi_mensuel(request)
        mois_filtre = request.query_params.get("mois") or ""
        mois_affiches = [m for m in grille["mois"] if not mois_filtre or m == mois_filtre]
        mois_courts = ["janv.", "févr.", "mars", "avr.", "mai", "juin", "juil.", "août", "sept.", "oct.", "nov.", "déc."]
        mois_longs = [
            "janvier", "février", "mars", "avril", "mai", "juin",
            "juillet", "août", "septembre", "octobre", "novembre", "décembre",
        ]

        def libelle_mois(mois, noms):
            annee, numero = mois.split("-")
            return f"{noms[int(numero) - 1]} {annee}"

        # `statut` : « payes » / « non_payes » pour n'imprimer que les élèves à jour, ou ceux qui
        # doivent encore quelque chose (même règle que l'écran — voir `_situation_paiement`).
        statut_filtre = request.query_params.get("statut") or ""
        titres_statut = {"payes": "Élèves payés", "non_payes": "Élèves non payés"}
        lignes = []
        total_paye = total_reste = Decimal("0")
        for e in grille["eleves"]:
            if statut_filtre in titres_statut and _situation_paiement(e, mois_filtre) != statut_filtre:
                continue
            par_mois = {m["mois"]: m for m in e["mois"]}
            # Total payé / reste sur les colonnes imprimées (inscription + mois, ou le mois filtré).
            colonnes = [par_mois[m] for m in mois_affiches if m in par_mois]
            if not mois_filtre and e["inscription"]:
                colonnes.append(e["inscription"])
            paye = sum((c["montant_paye"] for c in colonnes), Decimal("0"))
            reste = sum((c["reste"] for c in colonnes), Decimal("0"))
            total_paye += paye
            total_reste += reste
            lignes.append({
                "total_paye": paye, "reste": reste,
                "eleve_nom": e["eleve_nom"], "classe_nom": e["classe_nom"],
                "inscription": e["inscription"],
                "nb_impayes": sum(1 for m in e["mois"] if m["statut"] == "non_paye" and not m["a_venir"]),
                "nb_a_venir": sum(1 for m in e["mois"] if m["statut"] == "non_paye" and m["a_venir"]),
                "nb_partiels": sum(1 for m in e["mois"] if m["statut"] == "partiel"),
                "inscription_due": bool(e["inscription"]) and e["inscription"]["statut"] != "paye",
                "mois": [par_mois.get(m) for m in mois_affiches],
            })

        ecole = request.user.ecole
        html = render_to_string("payments/suivi_mensuel_classe_pdf.html", {
            "titre": grille["classe"], "annee_scolaire": grille["annee_scolaire"],
            "mois_filtre": libelle_mois(mois_filtre, mois_longs) if mois_filtre else "",
            "statut_filtre": titres_statut.get(statut_filtre, ""),
            "entetes_mois": [libelle_mois(m, mois_courts) for m in mois_affiches],
            "avec_inscription": not mois_filtre,
            "lignes": lignes, "total": len(lignes), "total_paye": total_paye, "total_reste": total_reste,
            "nb_colonnes_mois": len(mois_affiches) + (0 if mois_filtre else 1),
            "date_edition": timezone.now(),
            "ecole_nom": ecole.nom if ecole else "Taly-School",
            "ecole_logo_data_uri": _image_data_uri(ecole.logo, _mm_px(15, 15), mode="contain") if ecole and ecole.logo else None,
            "couleur_principale": ecole.couleur_principale if ecole else "#14304f",
        })
        buffer = BytesIO()
        pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        suffixe = f"_{statut_filtre}" if statut_filtre in titres_statut else ""
        response["Content-Disposition"] = f'attachment; filename="suivi_mensuel_{grille["annee_scolaire"]}{suffixe}.pdf"'
        return response

    @action(detail=False, methods=["post"], url_path="generer-pour-classe")
    def generer_pour_classe(self, request):
        """Génère en masse les frais de tous les élèves actifs d'une classe, pour une ou plusieurs
        types de frais (tous ceux de l'école si non précisé), sur une année scolaire — en reprenant
        le tarif paramétré pour cette classe (`TarifClasse`) si un existe, sinon le montant standard
        du type de frais. N'écrase jamais un frais déjà existant pour (élève, type de frais, année) —
        ou, pour un type Mensuel/Autre, pour le même mois d'échéance.

        Pour un type de frais mensuel, un élève exonéré de mensualité (Fondation gratuite, ou
        inscription/réinscription seulement — `EleveProfile.facteur_mensualite == 0`) n'a aucun
        frais généré pour ce type : il ne doit rien, inutile de créer une dette fictive qu'il
        faudrait ensuite suivre manuellement. La réduction « Fondation 50% »/fidélité, elle, n'est
        pas figée dans le montant du frais — elle est recalculée à la volée dans le suivi mensuel
        (`_calculer_suivi_mensuel`), pour rester à jour même si la catégorie de l'élève change
        après coup."""
        from academics.models import AnneeScolaire, Classe
        from people.models import EleveProfile

        classe_id = request.data.get("classe")
        annee_id = request.data.get("annee_scolaire")
        date_echeance = request.data.get("date_echeance")
        type_frais_ids = request.data.get("types_frais")  # optionnel : liste d'ids, sinon tous
        if not (classe_id and annee_id and date_echeance):
            raise ValidationError("Les champs 'classe', 'annee_scolaire' et 'date_echeance' sont requis.")

        classe = get_object_or_404(Classe, pk=classe_id, annee_scolaire__ecole_id=request.user.ecole_id)
        annee = get_object_or_404(AnneeScolaire, pk=annee_id, ecole_id=request.user.ecole_id)

        types_qs = TypeFrais.objects.filter(ecole_id=request.user.ecole_id)
        if type_frais_ids:
            types_qs = types_qs.filter(id__in=type_frais_ids)
        types_frais = list(types_qs)
        if not types_frais:
            raise ValidationError("Aucun type de frais à générer.")

        tarifs = {
            t.type_frais_id: t.montant
            for t in TarifClasse.objects.filter(classe=classe, annee_scolaire=annee, type_frais__in=types_frais)
        }

        try:
            echeance = date.fromisoformat(str(date_echeance))
        except ValueError:
            raise ValidationError("'date_echeance' doit être une date au format AAAA-MM-JJ.")

        eleves = list(EleveProfile.objects.filter(classe=classe, actif=True))
        # Mensuel / Autre : un frais par mois d'échéance (voir Frais.filtre_equivalents) — seul un
        # frais du MÊME mois bloque la génération, sinon le mois suivant ne pouvait jamais être
        # généré (ni donc encaissé). Tranche / Annuel : un seul frais de ce type pour l'année.
        par_mois = {TypeFrais.Periodicite.MENSUEL, TypeFrais.Periodicite.AUTRE}
        existants_qs = Frais.objects.filter(eleve__in=eleves, annee_scolaire=annee, type_frais__in=types_frais)
        existants = set(
            existants_qs.exclude(type_frais__periodicite__in=par_mois).values_list("eleve_id", "type_frais_id")
        ) | set(
            existants_qs.filter(Frais.filtre_mois(echeance), type_frais__periodicite__in=par_mois).values_list("eleve_id", "type_frais_id")
        )
        # Une seule inscription/réinscription par élève et par année, tous types confondus (voir
        # Frais.filtre_equivalents) — y compris entre deux types générés dans ce même appel.
        deja_inscrits = set(
            Frais.objects.filter(eleve__in=eleves, annee_scolaire=annee, type_frais__usage__in=USAGES_INSCRIPTION)
            .values_list("eleve_id", flat=True)
        )

        # Formule de scolarité déjà commencée par chaque élève (voir Frais.formule_scolarite) :
        # aucun frais d'une autre formule ne lui est généré pour cette année.
        formule_par_eleve = {}
        for eleve_id, periodicite in (
            Paiement.objects.filter(
                frais__eleve__in=eleves, frais__annee_scolaire=annee,
                frais__type_frais__periodicite__in=PERIODICITES_SCOLARITE,
            ).exclude(frais__type_frais__usage__in=USAGES_INSCRIPTION)
            .order_by("date_paiement", "id").values_list("frais__eleve_id", "frais__type_frais__periodicite")
        ):
            formule_par_eleve.setdefault(eleve_id, periodicite)

        a_creer = []
        for eleve in eleves:
            for type_frais in types_frais:
                formule = formule_par_eleve.get(eleve.id)
                if (
                    formule and type_frais.usage not in USAGES_INSCRIPTION
                    and type_frais.periodicite in PERIODICITES_SCOLARITE and type_frais.periodicite != formule
                ):
                    continue
                if (eleve.id, type_frais.id) in existants:
                    continue
                if eleve.exonere_fratrie and est_scolarite(type_frais):
                    continue  # « Élève Bonus » : pas de scolarité (voir people/fratrie.py)
                if type_frais.est_mensuel and (eleve.facteur_mensualite == 0 or echeance.month not in MOIS_MENSUALITE):
                    # Exonéré de mensualité, ou mois hors mensualité (Septembre : inscription ;
                    # Juillet/Août : hors année scolaire — voir MOIS_MENSUALITE).
                    continue
                if type_frais.usage in USAGES_INSCRIPTION:
                    if eleve.id in deja_inscrits:
                        continue
                    deja_inscrits.add(eleve.id)
                a_creer.append(Frais(
                    eleve=eleve, type_frais=type_frais, annee_scolaire=annee,
                    montant=tarifs.get(type_frais.id, type_frais.montant_standard), date_echeance=date_echeance,
                    mois=echeance.replace(day=1) if type_frais.periodicite in par_mois else None,
                ))
        Frais.objects.bulk_create(a_creer)
        return Response({"crees": len(a_creer), "classe": classe.nom, "eleves": len(eleves)})

    @action(detail=True, methods=["get"], url_path="fiche-paiement")
    def fiche_paiement(self, request, pk=None):
        """Fiche de paiement (PDF) d'un élève pour ce frais — accessible à l'élève/parent concerné."""
        frais = self.get_object()
        if frais.montant_paye <= 0:
            raise ValidationError("Aucun paiement n'a encore été enregistré pour cet élève sur ce frais.")
        pdf = _render_fiches_pdf([_fiche_context(frais)], "Fiche de paiement", ecole=frais.eleve.user.ecole)
        response = HttpResponse(pdf, content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="fiche_paiement_{frais.eleve.matricule}.pdf"'
        return response

    @action(detail=False, methods=["get"], url_path="fiches-paiement")
    def fiches_paiement(self, request):
        """Génère en un seul PDF les fiches de paiement de tous les élèves ayant déjà payé
        (respecte les mêmes filtres que la liste : classe, type de frais, année...)."""
        queryset = self.filter_queryset(self.get_queryset())
        fiches = [_fiche_context(f) for f in queryset if f.montant_paye > 0]
        if not fiches:
            raise ValidationError("Aucun élève n'a encore effectué de paiement pour cette sélection.")
        pdf = _render_fiches_pdf(fiches, "Fiches de paiement", ecole=request.user.ecole)
        response = HttpResponse(pdf, content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="fiches_paiement.pdf"'
        return response

    @action(detail=False, methods=["get"], url_path="proforma")
    def proforma(self, request):
        """Facture proforma (PDF) : récapitule tous les frais d'un élève pour une année scolaire
        (montant dû, déjà payé, solde) — un document d'estimation à présenter avant paiement,
        distinct de la « fiche de paiement » (qui documente des paiements déjà encaissés)."""
        from academics.models import AnneeScolaire
        from people.models import EleveProfile

        eleve_id = request.query_params.get("eleve")
        if not eleve_id:
            raise ValidationError("Le paramètre 'eleve' est requis.")
        eleve = get_object_or_404(EleveProfile, pk=eleve_id, user__ecole_id=request.user.ecole_id)

        if request.user.role == "student" and getattr(request.user, "eleve_profile", None) != eleve:
            raise ValidationError("Vous ne pouvez consulter que votre propre facture proforma.")
        if request.user.role == "parent" and eleve.parent_id != request.user.id:
            raise ValidationError("Vous ne pouvez consulter que la facture proforma de vos enfants.")

        annee_id = request.query_params.get("annee_scolaire")
        if annee_id:
            annee = get_object_or_404(AnneeScolaire, pk=annee_id, ecole_id=request.user.ecole_id)
        else:
            annee = annee_courante(request)
        if not annee:
            raise ValidationError("Aucune année scolaire active pour votre établissement.")

        lignes = list(
            Frais.objects.filter(eleve=eleve, annee_scolaire=annee).select_related("type_frais").order_by("date_echeance")
        )
        if not lignes:
            raise ValidationError("Aucun frais enregistré pour cet élève sur cette année scolaire.")

        total_du = sum((f.montant_du for f in lignes), Decimal("0"))
        total_paye = sum((f.montant_paye for f in lignes), Decimal("0"))

        ecole = eleve.user.ecole
        html = render_to_string("payments/proforma_pdf.html", {
            "eleve": eleve,
            "annee_scolaire": annee.libelle,
            "lignes": lignes,
            "total_du": total_du,
            "total_paye": total_paye,
            "total_solde": total_du - total_paye,
            "date_edition": date.today(),
            "ecole_nom": ecole.nom if ecole else "Taly-School",
            "ecole_adresse": ecole.adresse if ecole else "",
            "ecole_telephone": ecole.telephone if ecole else "",
            "ecole_logo_data_uri": _image_data_uri(ecole.logo, _mm_px(15, 15), mode="contain") if ecole else None,
            "couleur_principale": ecole.couleur_principale if ecole else "#14304f",
            "couleur_secondaire": ecole.couleur_secondaire if ecole else "#b8860b",
        })
        buffer = BytesIO()
        pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="proforma_{eleve.matricule}_{annee.libelle}.pdf"'
        return response

    @action(detail=False, methods=["get"], url_path="suivi-mensuel-pdf")
    def suivi_mensuel_pdf(self, request):
        """Version imprimable (PDF) du suivi mensuel des paiements de scolarité d'un élève —
        même calcul/permissions que `suivi_mensuel`, en document téléchargeable."""
        from academics.models import AnneeScolaire
        from people.models import EleveProfile

        eleve_id = request.query_params.get("eleve")
        if not eleve_id:
            raise ValidationError("Le paramètre 'eleve' est requis.")
        eleve = get_object_or_404(EleveProfile, pk=eleve_id, user__ecole_id=request.user.ecole_id)

        if request.user.role == "student" and getattr(request.user, "eleve_profile", None) != eleve:
            raise ValidationError("Vous ne pouvez consulter que votre propre suivi de paiement.")
        if request.user.role == "parent" and eleve.parent_id != request.user.id:
            raise ValidationError("Vous ne pouvez consulter que le suivi de paiement de vos enfants.")

        annee_id = request.query_params.get("annee_scolaire")
        if annee_id:
            annee = get_object_or_404(AnneeScolaire, pk=annee_id, ecole_id=request.user.ecole_id)
        else:
            annee = annee_courante(request)
        if not annee:
            raise ValidationError("Aucune année scolaire active pour votre établissement.")

        mois = _calculer_suivi_mensuel(eleve, annee)
        total_du = sum((m["montant_du"] for m in mois), Decimal("0"))
        total_paye = sum((m["montant_paye"] for m in mois), Decimal("0"))

        ecole = eleve.user.ecole
        html = render_to_string("payments/suivi_mensuel_pdf.html", {
            "eleve": eleve,
            "annee_scolaire": annee.libelle,
            "mois": mois,
            "total_du": total_du,
            "total_paye": total_paye,
            "total_solde": total_du - total_paye,
            "date_edition": date.today(),
            "ecole_nom": ecole.nom if ecole else "Taly-School",
            "ecole_adresse": ecole.adresse if ecole else "",
            "ecole_telephone": ecole.telephone if ecole else "",
            "ecole_logo_data_uri": _image_data_uri(ecole.logo, _mm_px(15, 15), mode="contain") if ecole and ecole.logo else None,
            "couleur_principale": ecole.couleur_principale if ecole else "#14304f",
            "couleur_secondaire": ecole.couleur_secondaire if ecole else "#b8860b",
        })
        buffer = BytesIO()
        pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
        response = HttpResponse(buffer.getvalue(), content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="suivi_mensuel_{eleve.matricule}_{annee.libelle}.pdf"'
        return response

    @action(detail=False, methods=["get"], url_path="summary")
    def summary(self, request):
        qs = self.get_queryset()
        # Totaux d'une année scolaire : `?annee_scolaire=` explicite, sinon l'année affichée
        # (`?toutes_annees=1` : toutes années confondues).
        if request.query_params.get("annee_scolaire"):
            qs = qs.filter(annee_scolaire_id=request.query_params["annee_scolaire"])
        else:
            qs = AnneeScolaireFilterBackend().filter_queryset(request, qs, self)
        # Sum("montant") sur la requête ne peut pas tenir compte de la réduction (catégorie de
        # paiement/fidélité) : Frais.montant_du n'est pas une colonne mais calculé par élève —
        # boucle Python, comme le fait déjà impayes_par_classe() ci-dessous pour la même raison.
        total_attendu = sum((f.montant_du for f in qs.select_related("eleve", "type_frais")), Decimal("0"))
        total_encaisse = Paiement.objects.filter(frais__in=qs).aggregate(total=Sum("montant"))["total"] or Decimal("0")
        taux = round(float(total_encaisse) / float(total_attendu) * 100, 1) if total_attendu else None
        return Response({
            "total_attendu": total_attendu,
            "total_encaisse": total_encaisse,
            "solde_total": total_attendu - total_encaisse,
            "taux_recouvrement": taux,
        })

    @action(detail=False, methods=["get"], url_path="impayes-par-classe")
    def impayes_par_classe(self, request):
        """Regroupe par classe tous les élèves ayant un solde restant à payer (respecte les
        mêmes filtres que la liste, ex: ?annee_scolaire=…) — et, à côté, ceux qui sont À JOUR
        (au moins un frais, tout réglé) : le comptable/l'administrateur voient les deux d'un
        coup d'œil plutôt que devoir déduire les « bons payeurs » par soustraction."""
        queryset = self.filter_queryset(self.get_queryset()).select_related("eleve__user", "eleve__classe", "type_frais")

        par_eleve: dict[int, dict] = {}
        for frais in queryset:
            entry = par_eleve.setdefault(frais.eleve_id, {
                "eleve": frais.eleve, "du": Decimal("0"), "paye": Decimal("0"),
            })
            entry["du"] += frais.montant_du
            entry["paye"] += frais.montant_paye

        par_classe: dict[int, dict] = {}
        for data in par_eleve.values():
            solde = data["du"] - data["paye"]
            eleve = data["eleve"]
            classe = eleve.classe
            classe_id = classe.id if classe else 0
            groupe = par_classe.setdefault(classe_id, {
                "classe_id": classe_id if classe else None,
                "classe_nom": classe.nom if classe else "Sans classe",
                "eleves": [], "eleves_a_jour": [],
            })
            fiche = {
                "eleve_id": eleve.id, "matricule": eleve.matricule,
                "nom_complet": eleve.user.get_full_name(), "du": data["du"],
                "paye": data["paye"], "solde": solde,
            }
            if solde > 0:
                groupe["eleves"].append(fiche)
            else:
                groupe["eleves_a_jour"].append(fiche)

        resultat = []
        for groupe in par_classe.values():
            groupe["eleves"].sort(key=lambda e: e["nom_complet"])
            groupe["eleves_a_jour"].sort(key=lambda e: e["nom_complet"])
            groupe["nb_impayes"] = len(groupe["eleves"])
            groupe["nb_a_jour"] = len(groupe["eleves_a_jour"])
            groupe["total_solde"] = sum((e["solde"] for e in groupe["eleves"]), Decimal("0"))
            resultat.append(groupe)
        # Une classe entièrement à jour (aucun impayé) reste incluse — comme avant, on n'exclut
        # que les classes sans aucun frais suivi du tout (ni impayé, ni à jour).
        resultat = [g for g in resultat if g["nb_impayes"] or g["nb_a_jour"]]
        resultat.sort(key=lambda g: g["classe_nom"])
        return Response(resultat)

    @action(detail=False, methods=["post"], url_path="notifier-impayes")
    def notifier_impayes(self, request):
        """Envoie immédiatement les rappels de paiement (SMS + e-mail) pour les frais en retard de
        l'école de l'utilisateur connecté — à chaque clic (`forcer=True`), même aux familles déjà
        relancées dans la semaine. Renvoie le compte-rendu détaillé (voir notifier_frais_impayes)."""
        from .notifications import notifier_frais_impayes
        return Response(notifier_frais_impayes(ecole_id=request.user.ecole_id, forcer=True))

    @action(detail=False, methods=["get"], url_path="export")
    def export(self, request):
        """Export CSV des frais/paiements (utilisable dans Excel)."""
        queryset = self.filter_queryset(self.get_queryset())
        response = HttpResponse(content_type="text/csv; charset=utf-8-sig")
        response["Content-Disposition"] = 'attachment; filename="frais_paiements.csv"'
        writer = csv.writer(response, delimiter=";")
        writer.writerow(["Élève", "Type de frais", "Tarif standard", "Montant dû", "Payé", "Solde", "Statut", "Échéance"])
        for frais in queryset:
            writer.writerow([
                frais.eleve.user.get_full_name(), frais.type_frais.nom, frais.montant, frais.montant_du,
                frais.montant_paye, frais.solde, frais.statut, frais.date_echeance,
            ])
        return response


class PaiementViewSet(viewsets.ModelViewSet):
    queryset = Paiement.objects.select_related("frais__eleve__user", "enregistre_par")
    serializer_class = PaiementSerializer
    permission_classes = [IsAdminOrComptabilite]
    filterset_fields = ["frais", "mode_paiement"]
    annee_scolaire_field = "frais__annee_scolaire"  # voir academics/annee.py

    def get_queryset(self):
        qs = super().get_queryset().filter(frais__eleve__user__ecole_id=self.request.user.ecole_id)
        user = self.request.user
        if user.role == "student" and hasattr(user, "eleve_profile"):
            return qs.filter(frais__eleve=user.eleve_profile)
        if user.role == "parent":
            return qs.filter(frais__eleve__parent=user)
        return qs

    def get_permissions(self):
        if self.action in ("list", "retrieve"):
            from rest_framework.permissions import IsAuthenticated
            return [IsAuthenticated()]
        if self.action == "destroy":
            # Suppression d'un paiement déjà encaissé : réservée à l'administrateur, même la
            # comptabilité (qui peut pourtant créer des paiements) n'y a pas accès — un paiement
            # supprimé par erreur fausse l'historique et le suivi mensuel de l'élève.
            return [IsAdmin()]
        return super().get_permissions()

    def create(self, request, *args, **kwargs):
        # Verrouille l'élève pendant validation + enregistrement : deux encaissements simultanés
        # du même mois (double clic, deux guichets) verraient sinon chacun le mois encore impayé
        # et passeraient tous les deux les garde-fous de PaiementSerializer.validate.
        from people.models import EleveProfile

        frais_id = str(request.data.get("frais", ""))
        with transaction.atomic():
            if frais_id.isdigit():
                list(EleveProfile.objects.select_for_update().filter(frais__pk=int(frais_id)).values_list("pk", flat=True))
            return super().create(request, *args, **kwargs)

    def perform_create(self, serializer):
        from accounts.models import JournalUtilisateur
        from accounts.services import journaliser

        paiement = serializer.save(enregistre_par=self.request.user)

        # Premier paiement sur ce frais : on fige la réduction en vigueur à cet instant (voir
        # Frais.facteur_applique/montant_du) — un changement ultérieur de catégorie de paiement
        # de l'élève n'affectera plus ce frais, la réduction ne s'applique jamais à une somme
        # déjà versée. Concerne la mensualité ET les frais d'inscription/réinscription (voir
        # Frais._facteur_reduction_courant) — `None` pour tout le reste (cantine, transport...).
        frais = paiement.frais
        facteur_actuel = frais._facteur_reduction_courant()
        if frais.facteur_applique is None and facteur_actuel is not None:
            frais.facteur_applique = facteur_actuel
            frais.save(update_fields=["facteur_applique"])

        journaliser(
            self.request.user, JournalUtilisateur.Categorie.PAIEMENT,
            f"Paiement enregistré : {paiement.frais.eleve.user.get_full_name()} — {paiement.montant} GNF",
            self.request,
        )

    def perform_destroy(self, instance):
        # Symétrique de perform_create() : si ce paiement était le dernier encaissé sur ce frais,
        # on "dégèle" la réduction (le frais suit de nouveau la catégorie courante de l'élève,
        # comme s'il n'avait jamais été payé) plutôt que de garder un gel devenu sans objet.
        frais = instance.frais
        instance.delete()
        if not frais.paiements.exists() and frais.facteur_applique is not None:
            frais.facteur_applique = None
            frais.save(update_fields=["facteur_applique"])
